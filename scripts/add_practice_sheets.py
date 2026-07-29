#!/usr/bin/env python3
"""
סקריפט חד-פעמי: (1) ממיר את עמודת "בניין" בגיליון 3 הקיים ("פעלים מוטים
מאומתים") מספרות רומיות לספרות רגילות, (2) מוסיף גיליון 4 ("תרגול פעלים") -
תרגיל לתלמידים עם עמודת "הפועל" (הצורה המנוקדת) מלאה ושאר העמודות ריקות
למילוי, ו-(3) גיליון 5 ("פתרונות") - אותן שורות בדיוק עם כל העמודות מלאות.
שני הגיליונות המקוריים (אוצר מילים כללי, אוצר מילים - פעלים) אינם נגעים כלל.

שימוש: python3 scripts/add_practice_sheets.py
"""
import json
import openpyxl
from openpyxl.styles import Font, Alignment

SRC = "אוצר_מילים_ערבית.xlsx"
SHEET3 = "פעלים מוטים מאומתים"
SHEET4 = "תרגול פעלים"
SHEET5 = "פתרונות"

ROMAN_TO_DIGIT = {"I": "1", "II": "2", "III": "3", "IV": "4", "V": "5",
                  "VI": "6", "VII": "7", "VIII": "8", "IX": "9", "X": "10"}

with open("scripts/practice_rows.json", encoding="utf-8") as f:
    practice_rows = json.load(f)

wb = openpyxl.load_workbook(SRC)

# --- (1) גיליון 3: המרת בניין מספרות רומיות לספרות רגילות ---
ws3 = wb[SHEET3]
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
    cell = row[1]  # עמודה B - בניין
    if cell.value in ROMAN_TO_DIGIT:
        cell.value = ROMAN_TO_DIGIT[cell.value]

# --- (2) + (3): גיליונות 4 ו-5 ---
header_font = Font(name="Arial", bold=True)
data_font = Font(name="Arial", bold=False)
center = Alignment(horizontal="center")
right = Alignment(horizontal="right")

headers = ["הפועל", "שורש", "בניין", "זמן", "גוף", "תרגום"]

for sheet_name in (SHEET4, SHEET5):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

ws4 = wb.create_sheet(SHEET4)
ws4.sheet_view.rightToLeft = True
ws5 = wb.create_sheet(SHEET5)
ws5.sheet_view.rightToLeft = True

for ws in (ws4, ws5):
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.alignment = center

for row_idx, item in enumerate(practice_rows, start=2):
    # עמודה A - הפועל (מנוקד) - מלאה בשני הגיליונות (זהו הנתון של התרגיל)
    for ws in (ws4, ws5):
        c = ws.cell(row=row_idx, column=1, value=item["arabic"])
        c.font = data_font
        c.alignment = center

    # גיליון 4 (תרגול): שאר העמודות ריקות למילוי התלמיד
    for col_idx in range(2, 7):
        ws4.cell(row=row_idx, column=col_idx, value=None).font = data_font

    # גיליון 5 (פתרונות): כל העמודות מלאות
    ws5.cell(row=row_idx, column=2, value=item["root"]).font = data_font
    ws5.cell(row=row_idx, column=2).alignment = center
    ws5.cell(row=row_idx, column=3, value=item["binyanDigit"]).font = data_font
    ws5.cell(row=row_idx, column=3).alignment = center
    ws5.cell(row=row_idx, column=4, value=item["tense"]).font = data_font
    ws5.cell(row=row_idx, column=4).alignment = center
    ws5.cell(row=row_idx, column=5, value=item["person"]).font = data_font
    ws5.cell(row=row_idx, column=5).alignment = center
    cell6 = ws5.cell(row=row_idx, column=6, value=item["meaning"])
    cell6.font = data_font
    cell6.alignment = right

for ws in (ws4, ws5):
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 40

wb.save(SRC)
print(f"עודכן גיליון 3 (בניין -> ספרות), ונוספו גיליונות '{SHEET4}' ו-'{SHEET5}' עם {len(practice_rows)} שורות כל אחד.")
