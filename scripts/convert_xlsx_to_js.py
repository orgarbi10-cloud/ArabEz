#!/usr/bin/env python3
"""
סקריפט חד-פעמי: ממיר את אוצר_מילים_ערבית.xlsx לשלושה קבצי JS סטטיים
(data/vocabulary.js, data/verb-roots.js, data/verb-practice-set.js) לשימוש
ישיר באתר, ללא תלות בקריאת xlsx בדפדפן. לא רץ כחלק מהאתר עצמו.
"""
import json
import openpyxl

SRC = "אוצר_מילים_ערבית.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)

# --- גיליון 1: אוצר מילים כללי ---
# מאגר המילים הרשמי של משרד החינוך (כיתות ז'-י"ב), מחולק ל-6 "פרקים" לפי
# רמת קושי מצטברת (פרק א' = לשעבר כיתה ז', וכן הלאה) - לא לפי נושא כמו
# האוצ"מ הקודם. בתוך כל פרק המילים מחולקות ל-5 "חלקים" (א'-ה', כ-30-45 מילים
# כל אחד) לפי רמת קושי יחסית + הקשר משמעות, כדי שלא יהיו מאות מילים ללימוד
# בבת אחת - בלי לערבב מילים בין פרקים שונים. עמודות אופציונליות (תעתיק/ריבוי/
# עתיד/מין/מושא/מענה) עשויות להיות ריקות לחלק מהמילים; מטופלות בעדינות
# (None/מחרוזת ריקה) ולא כשגיאה.
ws = wb["אוצר מילים כללי"]
rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)]

def blank_to_none(v):
    v = str(v).strip() if v is not None else ""
    return v or None

chapter_titles = {}
part_titles = {}  # (chapter_num, part_num) -> part_title
vocabulary = []
for (chapter_num, chapter_title, part_num, part_title, arabic, hebrew, translit,
     arabic_voc, plural, verb_present, gender, transitive, response) in rows:
    chapter_titles.setdefault(chapter_num, chapter_title)
    part_titles.setdefault((chapter_num, part_num), part_title)
    word = {
        "chapter": chapter_num,
        "part": part_num,
        "arabic": arabic.strip(),
        "arabicVoc": blank_to_none(arabic_voc),
        "hebrew": hebrew.strip(),
        "translit": blank_to_none(translit),
    }
    plural = blank_to_none(plural)
    verb_present = blank_to_none(verb_present)
    gender = blank_to_none(gender)
    response = blank_to_none(response)
    if plural: word["plural"] = plural
    if verb_present: word["verbPresent"] = verb_present
    if gender: word["gender"] = gender
    if blank_to_none(transitive): word["transitive"] = True
    if response: word["response"] = response
    vocabulary.append(word)

assert len(vocabulary) == 1129, f"Expected 1129 words, got {len(vocabulary)}"
assert len(chapter_titles) == 6, f"Expected 6 chapters, got {len(chapter_titles)}"
assert len(part_titles) == 30, f"Expected 30 chapter/part combos, got {len(part_titles)}"

chapters = [
    {
        "num": cn,
        "title": chapter_titles[cn],
        "parts": [
            {"num": pn, "title": part_titles[(cn, pn)]}
            for pn in sorted(pn for (ccn, pn) in part_titles if ccn == cn)
        ],
    }
    for cn in sorted(chapter_titles)
]

# --- גיליון 2: אוצר מילים - פעלים ---
def is_arabic_word(s):
    return isinstance(s, str) and all("؀" <= c <= "ۿ" for c in s.strip())

# הקובץ המקורי נותן את תרגום השורש כשם פועל ("לכתוב"). לצורכי האתר אנחנו
# מציגים במקום זאת את הפועל בעבר גוף הוא ("כתב"), כדי שיתאים לצורת הבסיס
# שמוצגת בכל התרגילים (עבר הוא = תבנית הבסיס של הבניינים). המרה ידנית קבועה,
# לא תרגום אוטומטי - כל 40 הערכים נבדקו ידנית.
INFINITIVE_TO_PAST = {
    "לכתוב": "כתב",
    "להבין": "הבין",
    "לשמוע": "שמע",
    "לשתות": "שתה",
    "ללכת": "הלך",
    "לחזור": "חזר",
    "לצאת": "יצא",
    "להיכנס": "נכנס",
    "לשבת": "ישב",
    "להסתכל": "הסתכל",
    "לעבוד": "עבד",
    "לשחק": "שיחק",
    "ללמוד": "למד",
    "לפתוח": "פתח",
    "לסגור": "סגר",
    "לבקש": "ביקש",
    "לדעת / להכיר": "ידע / הכיר",
    "לשנן / לשמור": "שינן / שמר",
    "לשבור": "שבר",
    "לכבס / לשטוף": "כיבס / שטף",
    "ללבוש": "לבש",
    "לרכוב": "רכב",
    "לרדת": "ירד",
    "לעלות": "עלה",
    "לשרת": "שירת",
    "לצייר": "צייר",
    "לבשל": "בישל",
    "לנטוע": "נטע",
    "לשאת / לסחוב": "נשא / סחב",
    "לדחוף / לשלם": "דחף / שילם",
    "לאסוף": "אסף",
    "לחתוך": "חתך",
    "להכות": "הכה",
    "לשלוט / לשפוט": "שלט / שפט",
    "להצליח": "הצליח",
    "להיכשל": "נכשל",
    "לרקוד": "רקד",
    "לצחוק": "צחק",
    "לקפוץ": "קפץ",
    "לשחות": "שחה",
}

ws2 = wb["אוצר מילים - פעלים"]
rows2 = [r for r in ws2.iter_rows(min_row=2, values_only=True) if is_arabic_word(r[0])]

# תנועת ר2 (העין) בבניין 1 היא מידע לקסיקלי לכל שורש בנפרד - אי אפשר לגזור
# אותה מכלל (למשל كتب=a/u אבל فهم=i/a). לכן נקראת מעמודות E/F בגיליון המקור
# (לא מנוחשת בקוד המנוע). ראו js/verbEngine.js.
VALID_VOWELS = {"a", "i", "u"}

verb_roots = []
for arabic, hebrew, translit, topic, past_vowel, present_vowel in rows2:
    root_letters = [c for c in arabic.strip() if c not in ("ـ",)]
    assert len(root_letters) == 3, f"Root not triliteral: {arabic!r} -> {root_letters}"
    hebrew_infinitive = hebrew.strip()
    assert hebrew_infinitive in INFINITIVE_TO_PAST, f"Missing past-tense mapping for: {hebrew_infinitive!r}"
    past_vowel = str(past_vowel).strip()
    present_vowel = str(present_vowel).strip()
    assert past_vowel in VALID_VOWELS, f"pastVowel לא תקין עבור {arabic!r}: {past_vowel!r}"
    assert present_vowel in VALID_VOWELS, f"presentVowel לא תקין עבור {arabic!r}: {present_vowel!r}"
    verb_roots.append({
        "root": arabic.strip(),
        "r1": root_letters[0],
        "r2": root_letters[1],
        "r3": root_letters[2],
        "hebrew": INFINITIVE_TO_PAST[hebrew_infinitive],
        "translit": translit.strip(),
        "topic": topic.strip(),
        "pastVowel": past_vowel,
        "presentVowel": present_vowel,
    })

assert len(verb_roots) == 40, f"Expected 40 verb roots, got {len(verb_roots)}"

with open("data/vocabulary.js", "w", encoding="utf-8") as f:
    f.write("// קובץ נוצר אוטומטית מתוך אוצר_מילים_ערבית.xlsx על ידי scripts/convert_xlsx_to_js.py\n")
    f.write("// אל תערוך ידנית - ערוך את קובץ המקור והרץ מחדש את הסקריפט.\n")
    f.write("//\n")
    f.write("// מקור: מאגר המילים הרשמי של משרד החינוך (כיתות ז'-י\"ב, הפיקוח על הוראת\n")
    f.write("// ערבית, מאי 2015). VOCAB_CHAPTERS מחולק ל-6 \"פרקים\" (א'-ו') לפי רמת קושי\n")
    f.write("// מצטברת - לא לפי נושא. כל פרק מחולק בתוכו ל-5 \"חלקים\" (parts, א'-ה')\n")
    f.write("// לפי רמת קושי יחסית + הקשר/משמעות, כדי לא להעמיס מאות מילים בבת אחת -\n")
    f.write("// בלי לערבב מילים בין פרקים. VOCABULARY word.chapter/word.part מצביעים\n")
    f.write("// לתוך VOCAB_CHAPTERS[].parts[]. arabic הוא הכתיב הרגיל (ללא ניקוד, לשימוש פנימי\n")
    f.write("// בהשוואות בבוחן ובחיפוש בלבד); arabicVoc הוא הניקוד המלא כפי שסופק במקור\n")
    f.write("// (לא ניקוד ידני שלנו הפעם). translit הוא תעתיק עברי מכני אות-באות (ללא\n")
    f.write("// סימון תנועות קצרות) לפי טבלת התאמה קבועה - ראו scripts/README/היסטוריית\n")
    f.write("// השיחה לפרטי המיפוי; יכול להיות null עבור שדות ללא ערך במקור.\n")
    f.write("// שדות אופציונליים כשקיימים במקור: plural (צורת ריבוי), verbPresent (צורת\n")
    f.write("// עתיד/הווה של פועל), gender (מין דקדוקי), transitive (הפועל דורש מושא),\n")
    f.write("// response (ברכת המענה הנהוגה, למילות ברכה/נימוס).\n\n")
    f.write("const VOCAB_CHAPTERS = ")
    f.write(json.dumps(chapters, ensure_ascii=False, indent=2))
    f.write(";\n\n")
    f.write("const VOCABULARY = ")
    f.write(json.dumps(vocabulary, ensure_ascii=False, indent=2))
    f.write(";\n")

with open("data/verb-roots.js", "w", encoding="utf-8") as f:
    f.write("// קובץ נוצר אוטומטית מתוך אוצר_מילים_ערבית.xlsx על ידי scripts/convert_xlsx_to_js.py\n")
    f.write("// אל תערוך ידנית - ערוך את קובץ המקור והרץ מחדש את הסקריפט.\n\n")
    f.write("const VERB_ROOTS = ")
    f.write(json.dumps(verb_roots, ensure_ascii=False, indent=2))
    f.write(";\n")

# --- גיליון 5: פתרונות (סט התרגול הקבוע ל"זהה את הפועל" באתר) ---
DIGIT_TO_BINYAN_ID = {"1": "I", "2": "II", "3": "III", "4": "IV", "5": "V", "6": "VI", "7": "VII", "8": "VIII", "10": "X"}
TENSE_LABEL_TO_KEY = {"עבר": "past", "הווה/עתיד": "present", "ציווי": "imperative", "מצדר": "masdar"}
PERSON_LABEL_TO_ID = {
    "אני": "ana", "אתה": "anta", "את": "anti", "הוא": "huwa", "היא": "hiya", "אנחנו": "nahnu",
    "אתם/ן השניים": "antuma", "אתם": "antum", "אתן": "antunna",
    "הם השניים": "huma_m", "הן השתיים": "huma_f", "הם": "hum", "הן": "hunna",
}

ws5 = wb["פתרונות"]
practice_set = []
seen_combos = set()
for verb, root, binyan_digit, tense_label, person_label, translation in ws5.iter_rows(min_row=2, values_only=True):
    if verb is None:
        continue
    binyan_id = DIGIT_TO_BINYAN_ID[str(binyan_digit)]
    tense = TENSE_LABEL_TO_KEY[tense_label]
    person_id = None if person_label == "–" else PERSON_LABEL_TO_ID[person_label]
    combo = (root, binyan_id, tense, person_id)
    assert combo not in seen_combos, f"כפילות בגיליון 5: {combo}"
    seen_combos.add(combo)
    # לא שומרים ar/he - נגזרים מחדש בזמן ריצה ע"י VerbEngine.generateForm, כדי
    # למנוע כפל נתונים מול המנוע (מקור האמת היחיד לצורות המנוטות).
    practice_set.append({"root": root, "binyanId": binyan_id, "tense": tense, "personId": person_id})

assert len(practice_set) == 550, f"Expected 550 practice items, got {len(practice_set)}"

with open("data/verb-practice-set.js", "w", encoding="utf-8") as f:
    f.write("// קובץ נוצר אוטומטית מתוך אוצר_מילים_ערבית.xlsx (גיליון 'פתרונות') על ידי\n")
    f.write("// scripts/convert_xlsx_to_js.py. אל תערוך ידנית - ערוך את קובץ המקור והרץ\n")
    f.write("// מחדש את הסקריפט. כל פריט מכיל רק {root, binyanId, tense, personId}; הצורה\n")
    f.write("// המנוטה (ar/he) נגזרת מחדש ע\"י VerbEngine.generateForm בזמן ריצה.\n\n")
    f.write("const VERB_PRACTICE_SET = ")
    f.write(json.dumps(practice_set, ensure_ascii=False, indent=2))
    f.write(";\n")

print(f"OK: {len(vocabulary)} words in {len(chapters)} chapters, {len(verb_roots)} verb roots, {len(practice_set)} practice items.")
