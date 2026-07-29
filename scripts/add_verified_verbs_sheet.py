#!/usr/bin/env python3
"""
סקריפט חד-פעמי: מוסיף גיליון שלישי לקובץ אוצר_מילים_ערבית.xlsx הקיים,
בשם "פעלים מוטים מאומתים", עם צירופי שורש+בניין שאומתו במחקר (ראה
scripts/verified_combos.json). שני הגיליונות הקיימים אינם נגעים כלל -
הקובץ נטען בלי data_only כדי לשמר את מבנהו המקורי במלואו, ומתווסף לו
רק גיליון חדש.

בשלב זה עמודות "זמן" ו"גוף" נשארות ריקות בכוונה (ייעודן: הנטייה המלאה
לכל גוף/זמן, שתמולא בהמשך ע"י scripts/generate_verified_conjugations.js
כשיתבקש להחיל את זה באתר).

שימוש: python3 scripts/add_verified_verbs_sheet.py
"""
import json
import openpyxl
from openpyxl.styles import Font, Alignment

SRC = "אוצר_מילים_ערבית.xlsx"
SHEET_NAME = "פעלים מוטים מאומתים"

with open("scripts/verified_combos.json", encoding="utf-8") as f:
    combos = json.load(f)

wb = openpyxl.load_workbook(SRC)

if SHEET_NAME in wb.sheetnames:
    del wb[SHEET_NAME]

ws = wb.create_sheet(SHEET_NAME)
ws.sheet_view.rightToLeft = True

headers = ["שורש (ערבית)", "בניין", "זמן", "גוף", "תרגום"]
header_font = Font(name="Arial", bold=True)
data_font = Font(name="Arial", bold=False)
center = Alignment(horizontal="center")
right = Alignment(horizontal="right")

for col_idx, title in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=title)
    cell.font = header_font
    cell.alignment = center

for row_idx, combo in enumerate(combos, start=2):
    ws.cell(row=row_idx, column=1, value=combo["root"]).font = data_font
    ws.cell(row=row_idx, column=1).alignment = center
    ws.cell(row=row_idx, column=2, value=combo["binyan"]).font = data_font
    ws.cell(row=row_idx, column=2).alignment = center
    ws.cell(row=row_idx, column=3, value=None).font = data_font  # זמן - ריק לעת עתה
    ws.cell(row=row_idx, column=4, value=None).font = data_font  # גוף - ריק לעת עתה
    cell5 = ws.cell(row=row_idx, column=5, value=combo["meaning"])
    cell5.font = data_font
    cell5.alignment = right

ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 40

wb.save(SRC)
print(f"נוסף גיליון '{SHEET_NAME}' עם {len(combos)} שורות (שני הגיליונות הקיימים לא נגעו).")
