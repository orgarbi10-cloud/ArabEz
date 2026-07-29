#!/usr/bin/env python3
"""
סקריפט חד-פעמי: מוסיף 320 שורות נוספות (scripts/practice_rows_batch2.json,
מאומתות כלא-חופפות למה שכבר קיים) לסוף גיליון 4 ("תרגול פעלים") וגיליון 5
("פתרונות") הקיימים - בלי למחוק/לשכתב את 180 השורות שכבר שם. שני הגיליונות
המקוריים (אוצר מילים כללי, אוצר מילים - פעלים) ושאר האתר אינם נגעים כלל.

שימוש: python3 scripts/append_practice_batch2.py
"""
import json
import openpyxl
from openpyxl.styles import Font, Alignment

SRC = "אוצר_מילים_ערבית.xlsx"
SHEET4 = "תרגול פעלים"
SHEET5 = "פתרונות"

with open("scripts/practice_rows_batch2.json", encoding="utf-8") as f:
    batch2 = json.load(f)

wb = openpyxl.load_workbook(SRC)
ws4 = wb[SHEET4]
ws5 = wb[SHEET5]

data_font = Font(name="Arial", bold=False)
center = Alignment(horizontal="center")
right = Alignment(horizontal="right")

start_row = ws4.max_row + 1
assert ws5.max_row == ws4.max_row, "גיליונות 4 ו-5 לא מסונכרנים במספר שורות"

for i, item in enumerate(batch2):
    row_idx = start_row + i

    c1 = ws4.cell(row=row_idx, column=1, value=item["arabic"])
    c1.font = data_font
    c1.alignment = center
    for col_idx in range(2, 7):
        ws4.cell(row=row_idx, column=col_idx, value=None).font = data_font

    c1b = ws5.cell(row=row_idx, column=1, value=item["arabic"])
    c1b.font = data_font
    c1b.alignment = center
    ws5.cell(row=row_idx, column=2, value=item["root"]).font = data_font
    ws5.cell(row=row_idx, column=2).alignment = center
    ws5.cell(row=row_idx, column=3, value=item["binyanDigit"]).font = data_font
    ws5.cell(row=row_idx, column=3).alignment = center
    ws5.cell(row=row_idx, column=4, value=item["tense"]).font = data_font
    ws5.cell(row=row_idx, column=4).alignment = center
    ws5.cell(row=row_idx, column=5, value=item["person"]).font = data_font
    ws5.cell(row=row_idx, column=5).alignment = center
    c6 = ws5.cell(row=row_idx, column=6, value=item["meaning"])
    c6.font = data_font
    c6.alignment = right

wb.save(SRC)
print(f"נוספו {len(batch2)} שורות לגיליונות '{SHEET4}' ו-'{SHEET5}' (סה\"כ כעת {ws4.max_row - 1} שורות תרגיל).")
